"""GAP-162: 成果物のクライアント共有リンク + 出力形式 (HTML/PDF/Excel)。

経営者質問 (2026-08-19):
  「あとは PDF や資料、エクセル、html など様々で出せる感じだよね？？」
  「これをこのままリンクとして資料を渡せる状態にもなっている？？」

設計:
  - 共有リンク: ランダム 32byte トークン。**DB にはハッシュのみ**保存するため、
    DB が漏れてもリンクは復元できない。期限つき + いつでも失効。
    閲覧は認証不要 (クライアントに渡すため) で、期限切れ・失効は 410 で断る。
  - 出力形式:
      HTML  … デザインテンプレ適用済みの実体をそのまま (印刷で PDF 化も可能)
      PDF   … 共有ページの「PDF で保存」= ブラウザ印刷 (A4 前提の CSS を活かす)
      Excel … HTML 内の表を xlsx へ変換 (見積・請求のように表が主役の成果物)。
              表が無い成果物は Excel 化できないと**正直に断る** (偽の空ファイルを出さない)。
"""

from __future__ import annotations

import hashlib
import io
import secrets
from dataclasses import dataclass
from datetime import UTC, datetime
from html.parser import HTMLParser
from typing import Any

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from src.audit import AuditEvent, AuditWriter

DEFAULT_EXPIRES_DAYS = 14
MAX_EXPIRES_DAYS = 180


class ShareError(Exception):
    """共有リンクの構造的失敗 (code: not_found / gone / no_html / no_table)。"""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True)
class ShareLink:
    id: str
    output_id: str
    label: str
    expires_at: datetime
    revoked_at: datetime | None
    view_count: int
    last_viewed_at: datetime | None
    created_at: datetime
    #: 発行直後のみ返る素のトークン (再取得は不可 — ハッシュしか保存しない)
    token: str | None = None


def _hash(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _row_to_link(row: Any, *, token: str | None = None) -> ShareLink:
    return ShareLink(
        id=str(row.id),
        output_id=str(row.output_id),
        label=str(row.label),
        expires_at=row.expires_at,
        revoked_at=row.revoked_at,
        view_count=int(row.view_count),
        last_viewed_at=row.last_viewed_at,
        created_at=row.created_at,
        token=token,
    )


_COLS = "id, output_id, label, expires_at, revoked_at, view_count, last_viewed_at, created_at"


async def create_share_link(
    session: AsyncSession,
    *,
    actor_id: str,
    output_id: str,
    label: str = "",
    expires_days: int = DEFAULT_EXPIRES_DAYS,
) -> ShareLink | None:
    """共有リンクを発行する。返り値 None = 成果物が不可視/不在。"""
    days = max(1, min(int(expires_days), MAX_EXPIRES_DAYS))
    visible = (
        await session.execute(
            text("select 1 from public.workflow_outputs where id = cast(:i as uuid)"),
            {"i": output_id},
        )
    ).first()
    if visible is None:
        return None
    token = secrets.token_urlsafe(32)
    row = (
        await session.execute(
            text(
                "insert into public.output_share_links "
                "(output_id, token_hash, label, expires_at, created_by) "
                "values (cast(:o as uuid), :h, :l, now() + make_interval(days => :d), "
                "        cast(:u as uuid)) "
                f"returning {_COLS}"
            ),
            {"o": output_id, "h": _hash(token), "l": label[:120], "d": days, "u": actor_id},
        )
    ).one()
    await AuditWriter(session).write(
        AuditEvent(
            action="output.share_link.create",
            target_type="workflow_output",
            actor_type="user",
            actor_id=actor_id,
            target_id=output_id,
            after={"expires_days": days, "label": label[:120]},
        )
    )
    return _row_to_link(row, token=token)


async def list_share_links(session: AsyncSession, *, output_id: str) -> list[ShareLink]:
    rows = (
        await session.execute(
            text(
                f"select {_COLS} from public.output_share_links "
                "where output_id = cast(:o as uuid) order by created_at desc"
            ),
            {"o": output_id},
        )
    ).all()
    return [_row_to_link(r) for r in rows]


async def revoke_share_link(
    session: AsyncSession, *, actor_id: str, link_id: str
) -> ShareLink | None:
    """リンクを失効させる (以後 410)。返り値 None = 不可視/不在。"""
    row = (
        await session.execute(
            text(
                "update public.output_share_links set revoked_at = now() "
                "where id = cast(:i as uuid) and revoked_at is null "
                f"returning {_COLS}"
            ),
            {"i": link_id},
        )
    ).first()
    if row is None:
        return None
    await AuditWriter(session).write(
        AuditEvent(
            action="output.share_link.revoke",
            target_type="workflow_output",
            actor_type="user",
            actor_id=actor_id,
            target_id=str(row.output_id),
            after={"link_id": link_id},
        )
    )
    return _row_to_link(row)


async def resolve_share_token(session: AsyncSession, *, token: str) -> tuple[str, str]:
    """共有トークンから (output_id, html) を返す。service セッション前提。

    期限切れ・失効は ShareError("gone")。存在しなければ ShareError("not_found")。
    """
    row = (
        await session.execute(
            text(
                "select id, output_id, expires_at, revoked_at "
                "from public.output_share_links where token_hash = :h"
            ),
            {"h": _hash(token)},
        )
    ).first()
    if row is None:
        raise ShareError("not_found", "共有リンクが見つかりません")
    if row.revoked_at is not None:
        raise ShareError("gone", "この共有リンクは無効化されています")
    if row.expires_at is not None and row.expires_at <= datetime.now(UTC):
        raise ShareError("gone", "この共有リンクは期限切れです")

    # GAP-319 (通し R3 所見 / G-14): 成果物そのものが生きていても、**案件やワークスペースが
    # 削除されていれば外部に開いたままにしない**。以前は workflow_outputs.deleted_at しか
    # 見ておらず、案件を削除しても WS を削除しても公開 URL が 200 で中身を返していた。
    out = (
        await session.execute(
            text(
                "select o.html_path from public.workflow_outputs o "
                "join public.projects p on p.id = o.project_id and p.deleted_at is null "
                "join public.workspaces w on w.id = p.workspace_id and w.deleted_at is null "
                "where o.id = :o and o.deleted_at is null"
            ),
            {"o": row.output_id},
        )
    ).first()
    if out is None or out.html_path is None:
        raise ShareError("no_html", "この成果物には共有できる内容がありません")
    html = await load_output_html(str(out.html_path))
    await session.execute(
        text(
            "update public.output_share_links "
            "set view_count = view_count + 1, last_viewed_at = now() where id = :i"
        ),
        {"i": row.id},
    )
    return str(row.output_id), html


async def load_output_html(storage_path: str) -> str:
    """成果物 HTML の実体を取る。mockdb は service 経路 (列 revoke 済のため)。"""
    from src.services.mocks.artifacts import (
        FILEDB_PREFIX,
        MOCKDB_PREFIX,
        fetch_content_service,
    )

    if storage_path.startswith(FILEDB_PREFIX):
        raise ShareError("no_html", "バイナリ形式の成果物はリンク共有に対応していません")
    if storage_path.startswith(MOCKDB_PREFIX):
        html = await fetch_content_service(storage_path[len(MOCKDB_PREFIX) :])
        if html is None:
            raise ShareError("no_html", "共有できる内容が見つかりません")
        return html
    from src.services.outputs.revise import download_html

    return await download_html(storage_path)


# ── 出力形式: HTML 内の表 → Excel ────────────────────────────────


class _TableParser(HTMLParser):
    """HTML の <table> を行列として取り出す (最初の 1 表ではなく全表)。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table: list[list[str]] | None = None
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        if tag == "table":
            self._table = []
        elif tag == "tr" and self._table is not None:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self._table is not None:
            if any(c for c in self._row):
                self._table.append(self._row)
            self._row = None
        elif tag == "table" and self._table is not None:
            if self._table:
                self.tables.append(self._table)
            self._table = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def html_tables_to_xlsx(html: str, *, title: str = "成果物") -> bytes:
    """HTML 内の表を xlsx にする。表が無ければ ShareError("no_table")。

    表が主役の成果物 (見積・請求・タスク一覧) を Excel で渡せるようにするための
    決定的変換。レイアウトまでは再現しない — 数字と項目を編集可能にするのが目的。
    """
    parser = _TableParser()
    parser.feed(html)
    if not parser.tables:
        raise ShareError(
            "no_table",
            "この成果物には表が無いため Excel 形式にできません (HTML / PDF をご利用ください)",
        )
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)  # pyright: ignore[reportArgumentType]
    for i, table_rows in enumerate(parser.tables, start=1):
        name = (title if i == 1 else f"{title}{i}")[:28] or f"表{i}"
        ws = wb.create_sheet(title=name)
        for row in table_rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def share_page_html(inner_html: str, *, title: str) -> str:
    """共有ページ: 成果物 HTML をそのまま見せ、印刷 (PDF 保存) 導線を足す。

    成果物のデザイン (テンプレ適用済み) を壊さないよう、上部バーは印刷時に消す。
    """
    bar = (
        '<div class="atelier-share-bar">'
        f"<span>{title}</span>"
        '<button type="button" onclick="window.print()">PDF で保存 / 印刷</button>'
        "</div>"
        "<style>"
        ".atelier-share-bar{position:sticky;top:0;z-index:9999;display:flex;"
        "align-items:center;justify-content:space-between;gap:12px;padding:8px 16px;"
        "background:#1F2937;color:#fff;font:600 13px/1.4 system-ui,sans-serif}"
        ".atelier-share-bar button{padding:6px 14px;border:0;border-radius:6px;"
        "background:#fff;color:#1F2937;font:600 12px/1 system-ui,sans-serif;cursor:pointer}"
        "@media print{.atelier-share-bar{display:none}}"
        "</style>"
    )
    if "<body" in inner_html:
        idx = inner_html.find("<body")
        end = inner_html.find(">", idx)
        if end >= 0:
            return inner_html[: end + 1] + bar + inner_html[end + 1 :]
    return bar + inner_html
